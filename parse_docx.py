import re
import docx
import fitz  # PyMuPDF
from bs4 import BeautifulSoup
from datetime import datetime, timedelta
import pandas as pd

def extract_data_from_table(docx_path):
    doc = docx.Document(docx_path)
    for table in doc.tables:
        first_cell_text = table.cell(0, 0).text
        if "УСЛУГИ ПО ПРЕДОСТАВЛЕНИЮ ДОСТУПА К ПРОГРАММНОМУ ПРОДУКТУ" in first_cell_text:
            right_cell_text = table.cell(3, 3).text
            dates = find_dates_in_section(right_cell_text)
            if dates:
                return dates
    return None

def find_dates_in_section(section):
    section = section.replace('\n', ' ')
    date_pattern = re.compile(r'\b\d{2}\.\d{2}\.\d{4}\b')
    dates = date_pattern.findall(section)
    return dates

def extract_text_from_pdf(pdf_path):
    document = fitz.open(pdf_path)
    text = ""
    for page_num in range(len(document)):
        page = document.load_page(page_num)
        text += page.get_text()
    return text

def find_section(text, start_pattern, end_pattern):
    pattern = re.compile(rf'({start_pattern}.*?){end_pattern}', re.DOTALL)
    match = pattern.search(text)
    if match:
        return match.group(0)
    return None

def find_payment_period(text):
    sections_to_check = ['2\\.6\\.2', '2\\.7\\.2']
    keywords = ['по факту', 'ежемесячно', 'ежеквартально']
    found_keywords = set()
    for section in sections_to_check:
        section_text = find_section(text, section, '\\n\\s*\\n')
        if section_text:
            for keyword in keywords:
                if keyword in section_text:
                    found_keywords.add(keyword)
    return found_keywords

def find_article_4_number(text):
    section = find_section(text, '4\\.1', '\\n\\s*\\н')
    if section:
        section = section.replace('4.1', '', 1)  # Убираем '4.1' один раз из начала
        match = re.search(r'\b(\d+)\b', section)
        if match:
            return int(match.group(1))  # Преобразуем в число
    return None

def extract_date_from_html(html_path):
    with open(html_path, 'r', encoding='utf-8') as file:
        soup = BeautifulSoup(file, 'html.parser')
        tables = soup.find_all('table')
        for table in tables:
            cell = table.find('td', string="Дата и время подписания ")
            if cell:
                rows = table.find_all('tr')
                if len(rows) >= 3:
                    columns = rows[2].find_all('td')
                    if len(columns) >= 4:
                        date_cell = columns[3]
                        # Оставляем только первые 10 символов
                        return date_cell.text[:10]
    return None

def parse_multiple_html(html_paths):
    html_dates = []
    for path in html_paths:
        date = extract_date_from_html(path)
        if date:
            html_dates.append(date)
    return html_dates

def get_first_working_day(year, month):
    # Получаем первый день месяца
    first_day = datetime(year, month, 1)
    # Проверяем, является ли первый день месяца рабочим днем
    while first_day.weekday() > 4:  # 0 - понедельник, 4 - пятница
        first_day += timedelta(days=1)
    return first_day

def add_business_days(start_date, business_days):
    current_date = start_date
    days_added = 0
    while days_added < business_days-1:
        current_date += timedelta(days=1)
        if current_date.weekday() < 5:  # 0 - понедельник, 4 - пятница
            days_added += 1
    return current_date

def calculate_deadline_dates(article_4_number, html_dates):
    deadlines = []
    for date_str in html_dates:
        year = int(date_str[-4:])
        month = int(date_str[3:5])
        first_working_day = get_first_working_day(year, month)
        deadline_date = add_business_days(first_working_day, article_4_number)
        deadlines.append(deadline_date.strftime('%d.%m.%Y'))
    return deadlines

def parse_docx_and_pdf(docx_path, pdf_path):
    docx_dates = extract_data_from_table(docx_path)
    if docx_dates:
        pdf_text = extract_text_from_pdf(pdf_path)
        payment_period = find_payment_period(pdf_text)
        article_4_number = find_article_4_number(pdf_text)
        return docx_dates, payment_period, article_4_number
    return None, None, None
