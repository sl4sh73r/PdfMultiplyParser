from docx import Document

def find_table_by_first_cell(doc, first_cell_text):
    for table in doc.tables:
        if table.cell(0, 0).text == first_cell_text:
            return table
    return None

def get_cell_value(table, row, col):
    return table.cell(row, col).text

def parse_dates_from_word(doc_path):
    # Открываем документ
    doc = Document(doc_path)
    
    # Текст первой ячейки таблицы, которую нужно найти
    first_cell_text = 'УСЛУГИ ПО ПРЕДОСТАВЛЕНИЮ ДОСТУПА К ПРОГРАММНОМУ ПРОДУКТУ'
    
    # Ищем таблицу по содержимому первой ячейки
    table = find_table_by_first_cell(doc, first_cell_text)
    
    if table:
        # Извлекаем данные из ячейки, содержащей срок (например, 4-я строка, 4-й столбец)
        row = 3  # Индексы начинаются с 0
        col = 3
        cell_value = get_cell_value(table, row, col).replace('\n', ' ')
        return cell_value
    else:
        print('Таблица с указанным содержимым первой ячейки не найдена.')
        return None